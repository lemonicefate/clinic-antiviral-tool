from dataclasses import replace
from datetime import date

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from app.cases import EXCLUDED, LotAllocation, READY_TO_EXPORT
from app.config import Settings
from app.database import Database
from app.his import InMemoryHisReader, Prescription
from app.main import create_app


def write_export_template(path):
    workbook = Workbook()
    codes = workbook.active
    codes.title = "OfficialCodes"
    codes.append(["ReasonCode", "ReasonName"])
    codes.append(["R01", "符合公費使用條件"])
    report = workbook.create_sheet("Report")
    report.append(
        [
            "source_id",
            "medical_record_number",
            "patient_name",
            "patient_id",
            "birth_date",
            "sex",
            "report_date",
            "doctor_id",
            "official_material",
            "reason_code",
            "reason_label",
            "actual_quantity",
            "lot_number",
            "lot_quantity",
        ]
    )
    workbook.save(path)


def make_prescription(source_id, *, quantity=10):
    return Prescription(
        source_id=source_id,
        medical_record_number=source_id,
        patient_name=f"Patient {source_id}",
        patient_id=source_id,
        birth_date=date(1980, 1, 2),
        sex="M",
        report_date=date(2026, 10, 9),
        doctor_id="D001",
        medication_id="eraflu",
        medication_name="易剋冒膠囊",
        official_material="DDMTR2018090002:易剋冒膠囊(顆)",
        unit="capsule",
        source_quantity=quantity,
        treat="Y",
    )


class MutableReader:
    def __init__(self, prescriptions):
        self.prescriptions = list(prescriptions)

    def _reader(self):
        return InMemoryHisReader(self.prescriptions)

    def most_recent(self, doctor_id):
        return self._reader().most_recent(doctor_id)

    def find_by_mrn(self, medical_record_number, doctor_id=None):
        return self._reader().find_by_mrn(medical_record_number, doctor_id)

    def scan(self, start_date, end_date):
        return self._reader().scan(start_date, end_date)

    def get(self, source_id):
        return self._reader().get(source_id)


def make_ready_case(database, item):
    case = database.ensure_case(item)
    return database.save_pharmacy_fields(
        case.id,
        reason_code="R01",
        reason_label="符合公費使用條件",
        actual_quantity=item.source_quantity,
        lots=(LotAllocation("LOT-A", item.source_quantity),),
        operator="P001",
    )


def make_client(tmp_path, items):
    settings = Settings.for_testing(tmp_path)
    write_export_template(settings.template_path)
    database = Database(settings.database_path)
    reader = MutableReader(items)
    client = TestClient(create_app(settings=settings, his_reader=reader, database=database))
    return client, database, reader


def test_export_skips_invalid_cases_and_creates_downloadable_batch(tmp_path):
    valid_item = make_prescription("VALID")
    invalid_item = make_prescription("INVALID")
    client, database, _ = make_client(tmp_path, [valid_item, invalid_item])
    valid = make_ready_case(database, valid_item)
    invalid = database.ensure_case(invalid_item)

    response = client.post(
        "/pharmacy/export",
        data={"export_case_ids": [str(valid.id), str(invalid.id)], "operator": "P001"},
    )

    assert response.status_code == 200
    assert "Selected: 2" in response.text
    assert "Exported: 1" in response.text
    assert "Skipped: 1" in response.text
    assert "missing medication-use reason" in response.text
    batch = database.get_export_batch(1)
    assert batch["filename"]
    assert client.get("/pharmacy/export/1/download").status_code == 200
    workbook = load_workbook(tmp_path / "app-data" / "exports" / batch["filename"], data_only=True)
    rows = list(workbook["Report"].iter_rows(min_row=2, values_only=True))
    assert rows[0][0] == valid.source_id
    assert all(row[0] != invalid.source_id for row in rows if row[0])


def test_changed_his_case_shows_changed_fields_and_keep_choice_allows_export(tmp_path):
    original = make_prescription("CHANGED")
    client, database, reader = make_client(tmp_path, [original])
    case = make_ready_case(database, original)
    reader.prescriptions = [replace(original, source_quantity=12)]

    blocked = client.post("/pharmacy/export", data={"export_case_ids": [str(case.id)], "operator": "P001"})

    assert "source quantity" in blocked.text
    assert "Use latest HIS data" in blocked.text
    assert "Keep current reporting data" in blocked.text
    assert "Exclude case" in blocked.text
    assert database.get_export_batch(1) is None

    resolved = client.post(
        "/pharmacy/export",
        data={
            "export_case_ids": [str(case.id)],
            "operator": "P001",
            f"resolution_{case.id}": "keep",
        },
    )

    assert "Exported: 1" in resolved.text
    saved = database.get_case(case.id)
    assert saved.source_fingerprint == reader.get(original.source_id).reporting_fingerprint
    assert saved.source_quantity == original.source_quantity


def test_changed_case_can_be_excluded_while_other_valid_case_exports(tmp_path):
    changed = make_prescription("CHANGED")
    unaffected = make_prescription("OK")
    client, database, reader = make_client(tmp_path, [changed, unaffected])
    changed_case = make_ready_case(database, changed)
    unaffected_case = make_ready_case(database, unaffected)
    reader.prescriptions = [replace(changed, source_quantity=12), unaffected]

    blocked = client.post(
        "/pharmacy/export",
        data={
            "export_case_ids": [str(changed_case.id), str(unaffected_case.id)],
            "operator": "P001",
        },
    )
    assert "source quantity" in blocked.text

    resolved = client.post(
        "/pharmacy/export",
        data={
            "export_case_ids": [str(changed_case.id), str(unaffected_case.id)],
            "operator": "P001",
            f"resolution_{changed_case.id}": "exclude",
        },
    )

    assert "Selected: 2" in resolved.text
    assert "Exported: 1" in resolved.text
    assert database.get_case(changed_case.id).state == EXCLUDED
    assert database.get_case(unaffected_case.id).state == READY_TO_EXPORT


def test_changed_case_can_use_latest_his_quantity_before_export(tmp_path):
    original = make_prescription("LATEST")
    client, database, reader = make_client(tmp_path, [original])
    case = make_ready_case(database, original)
    reader.prescriptions = [replace(original, source_quantity=12)]

    client.post("/pharmacy/export", data={"export_case_ids": [str(case.id)], "operator": "P001"})
    resolved = client.post(
        "/pharmacy/export",
        data={
            "export_case_ids": [str(case.id)],
            "operator": "P001",
            f"resolution_{case.id}": "latest",
        },
    )

    assert "Exported: 1" in resolved.text
    saved = database.get_case(case.id)
    assert saved.source_quantity == 12
    assert saved.actual_quantity == 12
    assert saved.lots[0].quantity == 12
