from openpyxl import Workbook

import pytest

from app.template import ReasonOption, TemplateCompatibilityError, load_reason_options


def write_reason_template(path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "OfficialCodes"
    sheet.append(["ReasonCode", "ReasonName"])
    sheet.append(["R01", "符合公費使用條件"])
    sheet.append(["R02", "其他正式條件"])
    workbook.save(path)


def test_reason_options_are_read_from_workbook_code_table(tmp_path):
    path = tmp_path / "smis_report.xlsx"
    write_reason_template(path)

    assert load_reason_options(path) == [
        ReasonOption(code="R01", label="符合公費使用條件"),
        ReasonOption(code="R02", label="其他正式條件"),
    ]


def test_incompatible_template_fails_with_stable_code(tmp_path):
    path = tmp_path / "wrong.xlsx"
    workbook = Workbook()
    workbook.active.append(["wrong", "columns"])
    workbook.save(path)

    with pytest.raises(TemplateCompatibilityError) as error:
        load_reason_options(path)

    assert error.value.code == "EXPORT_TEMPLATE_001"


def test_official_single_column_reasons_preserve_distinct_subconditions(tmp_path):
    path = tmp_path / "official.xlsx"
    workbook = Workbook()
    workbook.active.title = "住院病患"
    workbook.active.append(["姓名", "用藥對象"])
    codes = workbook.create_sheet("代碼表")
    codes.append(["性別", "物資項目代碼", "用藥對象"])
    values = ["5:合成測試條件--1:子條件甲", "5:合成測試條件--2:子條件乙", "23:合成測試條件"]
    for value in values:
        codes.append(["0:女", "TEST:合成藥品", value])
    workbook.save(path)

    assert load_reason_options(path) == [ReasonOption(code=value, label=value) for value in values]


@pytest.mark.parametrize("values", [["missing-code"], ["5:合成條件", "5:合成條件"], ["5:"]])
def test_official_code_table_rejects_malformed_or_duplicate_values(tmp_path, values):
    path = tmp_path / "official.xlsx"
    workbook = Workbook()
    workbook.active.title = "代碼表"
    workbook.active.append(["用藥對象"])
    for value in values:
        workbook.active.append([value])
    workbook.save(path)

    with pytest.raises(TemplateCompatibilityError):
        load_reason_options(path)
