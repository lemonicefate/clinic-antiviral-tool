from datetime import date

from openpyxl import Workbook, load_workbook
import pytest

from app.cases import LotAllocation, READY_TO_EXPORT
from app.database import CaseRecord
from app.exporter import ExportError, export_cases, find_export_table, validate_case_for_export
from app.template import TemplateCompatibilityError


def write_export_template(path):
    workbook = Workbook()
    codes = workbook.active
    codes.title = "OfficialCodes"
    codes.append(["ReasonCode", "ReasonName"])
    codes.append(["R01", "符合公費使用條件"])
    report = workbook.create_sheet("Report")
    report.append(
        [
            "source_id",
            "medical_record_number",
            "patient_name",
            "patient_id",
            "birth_date",
            "sex",
            "report_date",
            "doctor_id",
            "official_material",
            "reason_code",
            "reason_label",
            "actual_quantity",
            "lot_number",
            "lot_quantity",
        ]
    )
    workbook.save(path)


def make_case(case_id=1, *, state=READY_TO_EXPORT, lots=(LotAllocation("LOT-A", 10),)):
    return CaseRecord(
        id=case_id,
        source_id=f"SOURCE-{case_id}",
        medical_record_number=f"MRN-{case_id}",
        patient_name="王小明",
        patient_id="P001",
        birth_date=date(1980, 1, 2),
        sex="M",
        report_date=date(2026, 10, 9),
        doctor_id="D001",
        medication_id="eraflu",
        medication_name="易剋冒膠囊",
        official_material="DDMTR2018090002:易剋冒膠囊(顆)",
        source_quantity=10,
        actual_quantity=10,
        reason_code="R01",
        reason_label="符合公費使用條件",
        state=state,
        exclusion_reason=None,
        source_fingerprint=(),
        created_at="2026-10-09T00:00:00+00:00",
        updated_at="2026-10-09T00:00:00+00:00",
        lots=lots,
    )


def test_export_copies_template_writes_each_lot_and_reopens_output(tmp_path):
    template = tmp_path / "template.xlsx"
    output = tmp_path / "output.xlsx"
    write_export_template(template)
    case = make_case(lots=(LotAllocation("LOT-A", 6), LotAllocation("LOT-B", 4)))

    result = export_cases(template, output, [case])

    assert result.row_count == 2
    workbook = load_workbook(output, data_only=True)
    rows = list(workbook["Report"].iter_rows(min_row=2, values_only=True))
    assert rows[0][8] == case.official_material
    assert rows[0][9] == "R01"
    assert rows[0][11] == 10
    assert [(row[12], row[13]) for row in rows] == [("LOT-A", 6), ("LOT-B", 4)]


def test_export_validation_returns_concise_reasons_for_invalid_case():
    case = make_case(lots=(LotAllocation("LOT-A", 9),))
    invalid = CaseRecord(**{**case.__dict__, "reason_code": None, "actual_quantity": 10})

    reasons = validate_case_for_export(invalid)

    assert "missing medication-use reason" in reasons
    assert "lot quantities do not equal actual quantity" in reasons


def test_incompatible_export_template_fails_clearly(tmp_path):
    template = tmp_path / "wrong.xlsx"
    workbook = Workbook()
    workbook.active.append(["not", "an", "export", "table"])
    workbook.save(template)

    with pytest.raises(TemplateCompatibilityError) as error:
        export_cases(template, tmp_path / "output.xlsx", [make_case()])

    assert error.value.code == "EXPORT_TEMPLATE_001"
