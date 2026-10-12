from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook


class TemplateCompatibilityError(RuntimeError):
    code = "EXPORT_TEMPLATE_001"

    def __init__(self, message: str):
        super().__init__(message)
        self.message = message


@dataclass(frozen=True)
class ReasonOption:
    code: str
    label: str


CODE_HEADERS = {
    "code",
    "reasoncode",
    "reason_code",
    "代碼",
    "代码",
    "使用條件代碼",
    "使用条件代码",
}
LABEL_HEADERS = {
    "label",
    "name",
    "reason",
    "reasonname",
    "reason_name",
    "說明",
    "说明",
    "名稱",
    "名称",
    "使用條件",
    "使用条件",
    "用藥原因",
    "用药原因",
}


def load_reason_options(path: Path) -> list[ReasonOption]:
    if not path.exists() or not path.is_file():
        raise TemplateCompatibilityError("The official Excel template is missing.")
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as error:
        raise TemplateCompatibilityError("The official Excel template cannot be opened.") from error

    try:
        if "代碼表" in workbook.sheetnames:
            options = _official_reason_options(workbook["代碼表"].iter_rows(values_only=True))
            if options:
                return options
            raise TemplateCompatibilityError("The official Excel reason code table is missing or incompatible.")
        for sheet in workbook.worksheets:
            options = _options_from_sheet(sheet.iter_rows(values_only=True))
            if options:
                return options
    finally:
        workbook.close()
    raise TemplateCompatibilityError("The official Excel reason code table is missing or incompatible.")


def _official_reason_options(rows: Iterable[tuple[object, ...]]) -> list[ReasonOption]:
    rows = iter(rows)
    headers = next(rows, ())
    indexes = [index for index, value in enumerate(headers) if _header_name(value) == "用藥對象"]
    if len(indexes) != 1:
        return []
    index = indexes[0]
    options: list[ReasonOption] = []
    seen: set[str] = set()
    for row in rows:
        value = _text(row[index] if index < len(row) else "")
        if not value:
            continue
        code, separator, label = value.partition(":")
        if not separator or not code.isdecimal() or not label.strip() or value in seen:
            raise TemplateCompatibilityError("The official Excel reason code table contains invalid or duplicate values.")
        # The complete dropdown value identifies subconditions sharing a main code.
        options.append(ReasonOption(code=value, label=value))
        seen.add(value)
    return options


def _options_from_sheet(rows: Iterable[tuple[object, ...]]) -> list[ReasonOption]:
    rows = iter(rows)
    for row in rows:
        indexes = _header_indexes(row)
        if indexes is None:
            continue
        code_index, label_index = indexes
        options: list[ReasonOption] = []
        for value_row in rows:
            code = _text(value_row[code_index] if code_index < len(value_row) else "")
            label = _text(value_row[label_index] if label_index < len(value_row) else "")
            if not code and not label:
                if options:
                    break
                continue
            if code and label:
                options.append(ReasonOption(code=code, label=label))
        if options:
            return options
    return []


def _header_indexes(row: tuple[object, ...]) -> tuple[int, int] | None:
    code_index = next(
        (index for index, value in enumerate(row) if _header_name(value) in CODE_HEADERS),
        None,
    )
    label_index = next(
        (index for index, value in enumerate(row) if _header_name(value) in LABEL_HEADERS),
        None,
    )
    if code_index is None or label_index is None or code_index == label_index:
        return None
    return code_index, label_index


def _header_name(value: object) -> str:
    return "".join(character for character in _text(value).lower() if character.isalnum() or character in "_中英文用藥條件代碼說明名稱原因")


def _text(value: object) -> str:
    return str(value or "").strip()
